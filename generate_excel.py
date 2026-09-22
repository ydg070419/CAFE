import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_excel_specification():
    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------
    # 1. 시트: 주문서 항목 명세
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "주문서_항목_명세"
    ws1.views.sheetView[0].showGridLines = True

    # 타이틀
    ws1.merge_cells("A1:G1")
    title_cell = ws1["A1"]
    title_cell.value = "☕ 바이브 카페 주문서 - 입력 항목 명세서"
    title_cell.font = Font(name="맑은 고딕", size=16, bold=True, color="1F497D")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 40

    headers1 = ["순번", "입력 항목명", "HTML 태그/타입", "연결 라벨 (<label>)", "주요 속성 (Attributes)", "필수 여부", "항목 설명 및 기능"]
    ws1.append(headers1)
    ws1.row_dimensions[2].height = 28

    items_data = [
        [1, "이름", "<input type=\"text\">", "for=\"customer-name\"", "id=\"customer-name\"\nname=\"customerName\"\nplaceholder=\"예: 홍길동\"\nrequired", "필수 (required)", "주문 고객의 성함을 한 줄 텍스트로 입력받는 필드입니다. 미입력 시 폼 제출이 차단됩니다."],
        [2, "전화번호", "<input type=\"tel\">", "for=\"customer-phone\"", "id=\"customer-phone\"\nname=\"customerPhone\"\nplaceholder=\"예: 010-1234-5678\"", "선택", "연락처를 입력받는 필드로, 모바일 브라우저 접속 시 숫자 키패드를 유도하여 입력을 편리하게 돕습니다."],
        [3, "음료 선택", "<select>, <option>", "for=\"beverage-select\"", "id=\"beverage-select\"\nname=\"beverage\"", "선택", "드롭다운 목록으로 아메리카노, 카페라떼, 카페모카, 녹차라떼 중 하나의 음료를 선택할 수 있습니다."],
        [4, "사이즈 선택", "<input type=\"radio\">", "for=\"size-s\"\nfor=\"size-m\"\nfor=\"size-l\"", "name=\"cupSize\"\nvalue=\"S\", \"M\", \"L\"\nchecked (M 기본선택)", "선택 (택1)", "<fieldset> 및 <legend>로 묶인 단일 선택 라디오 버튼입니다. S(작은), M(보통), L(큰) 중 1개만 선택 가능합니다."],
        [5, "요청사항", "<textarea>", "for=\"special-requests\"", "id=\"special-requests\"\nname=\"specialRequests\"\nrows=\"4\"\ncols=\"45\"\nplaceholder=안내문구", "선택", "시럽 추가, 얼음 양 조절 등 추가 요구사항을 여러 줄로 자유롭게 입력받을 수 있는 텍스트 상자입니다."],
        [6, "주문하기", "<button type=\"submit\">", "-", "type=\"submit\"", "-", "작성한 폼 데이터를 지정된 서버 대상(action)으로 전송(제출)하는 버튼입니다."],
        [7, "다시 작성", "<button type=\"reset\">", "-", "type=\"reset\"", "-", "폼 양식에 입력되어 있던 모든 내용을 초기 상태로 한 번에 리셋(초기화)하는 버튼입니다."]
    ]

    for row in items_data:
        ws1.append(row)

    # -------------------------------------------------------------
    # 2. 시트: HTML 태그 가이드
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="HTML_태그_가이드")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:E1")
    title2 = ws2["A1"]
    title2.value = "📋 사용된 HTML5 태그 및 속성 가이드"
    title2.font = Font(name="맑은 고딕", size=16, bold=True, color="1F497D")
    title2.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 40

    headers2 = ["태그명", "분류", "주요 속성", "태그 역할 및 설명", "바이브 카페 적용 예시"]
    ws2.append(headers2)
    ws2.row_dimensions[2].height = 28

    tags_data = [
        ["<!DOCTYPE html>", "문서 선언", "-", "웹 브라우저에게 현재 문서가 HTML5 최신 표준 규격을 따르고 있음을 알림", "<!DOCTYPE html>"],
        ["<html lang=\"ko\">", "루트 요소", "lang", "HTML 문서의 최상위 태그이며, lang=\"ko\"는 기본 언어가 한국어임을 검색엔진과 스크린 리더에 알림", "<html lang=\"ko\">...</html>"],
        ["<head>", "문서 정보", "-", "브라우저 화면에는 직접 나타나지 않는 문서의 메타데이터(설정, 제목 등)를 포함", "<head><meta ...><title>...</title></head>"],
        ["<meta>", "메타데이터", "charset, name, content", "문자 인코딩(UTF-8) 및 반응형 뷰포트(viewport) 환경 설정", "<meta charset=\"UTF-8\">"],
        ["<title>", "페이지 제목", "-", "웹 브라우저 탭 상단 및 즐겨찾기(북마크)에 표시될 문서 제목 정의", "<title>☕ 바이브 카페 주문서</title>"],
        ["<body>", "본문 콘텐츠", "-", "웹 브라우저 화면에 실제로 사용자에게 렌더링되는 모든 시각적 요소를 포함", "<body><h1>...</h1><form>...</form></body>"],
        ["<h1>", "제목(Heading)", "-", "페이지의 최상위 대제목을 정의 (가장 큰 글씨 및 핵심 주제)", "<h1>☕ 바이브 카페 주문서</h1>"],
        ["<p>", "단락(Paragraph)", "-", "텍스트 문단을 구분하여 줄바꿈 및 상하 여백을 제공", "<p>바이브 카페에 오신 것을 환영합니다!</p>"],
        ["<hr>", "구분선", "-", "콘텐츠의 주제 변경이나 구분을 위한 시각적 수평 가로선을 표시", "<hr>"],
        ["<form>", "입력 폼", "action, method", "사용자가 입력한 데이터를 그룹화하여 서버로 전송하는 양식 컨테이너", "<form action=\"#\" method=\"post\">"],
        ["<label>", "라벨(이름표)", "for", "입력 폼 요소의 이름표로, for 속성을 input의 id와 연결해 클릭 접근성 향상", "<label for=\"customer-name\">이름</label>"],
        ["<input>", "데이터 입력창", "type, id, name, placeholder, required, checked, value", "text, tel, radio 등 다양한 유형의 입력 필드를 생성하는 가장 핵심적인 태그", "<input type=\"text\" placeholder=\"홍길동\" required>"],
        ["<select>", "드롭다운", "id, name", "여러 선택지 중 사용자가 하나를 고를 수 있는 펼침 목록 메뉴 생성", "<select id=\"beverage-select\" name=\"beverage\">"],
        ["<option>", "선택 항목", "value", "드롭다운 목록(<select>) 내부에 들어가는 개별 선택지 정의", "<option value=\"americano\">아메리카노</option>"],
        ["<fieldset>", "폼 그룹화", "-", "폼 요소 내에서 관련된 항목들을 시각적 테두리 상자로 묶어 정리", "<fieldset><legend>사이즈</legend>...</fieldset>"],
        ["<legend>", "그룹 캡션", "-", "<fieldset>으로 묶인 상자의 제목(캡션)을 정의", "<legend>4. 사이즈 선택</legend>"],
        ["<textarea>", "장문 입력", "id, name, rows, cols, placeholder", "한 줄 이상의 긴 텍스트나 요청사항을 여러 줄로 자유롭게 입력받는 텍스트 박스", "<textarea rows=\"4\" cols=\"45\"></textarea>"],
        ["<button>", "버튼", "type", "submit(제출), reset(초기화) 등 폼 제어 및 사용자 클릭 동작 수행", "<button type=\"submit\">주문하기</button>"]
    ]

    for row in tags_data:
        ws2.append(row)

    # -------------------------------------------------------------
    # 3. 시트: 소스 코드
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="주문서_HTML_소스코드")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:C1")
    title3 = ws3["A1"]
    title3.value = "📄 바이브 카페 주문서 (index.html) 소스코드 전문"
    title3.font = Font(name="맑은 고딕", size=16, bold=True, color="1F497D")
    title3.alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[1].height = 40

    headers3 = ["줄 번호", "코드 라인", "비고"]
    ws3.append(headers3)
    ws3.row_dimensions[2].height = 28

    with open(r"c:\Users\dongg\OneDrive\Desktop\AI개발과실전\4주차\카페프로젝트\index.html", "r", encoding="utf-8") as f:
        lines = f.readlines()

    for idx, line in enumerate(lines, start=1):
        clean_line = line.rstrip("\r\n")
        note = ""
        if "DOCTYPE" in clean_line:
            note = "HTML5 표준 선언"
        elif "<form" in clean_line:
            note = "주문서 폼 시작"
        elif "customer-name" in clean_line and "input" in clean_line:
            note = "1. 이름 입력 필드"
        elif "customer-phone" in clean_line and "input" in clean_line:
            note = "2. 전화번호 입력 필드"
        elif "<select" in clean_line:
            note = "3. 음료 드롭다운 선택 필드"
        elif "<fieldset" in clean_line:
            note = "4. 사이즈 선택 그룹 시작"
        elif "<textarea" in clean_line:
            note = "5. 요청사항 여러 줄 입력 필드"
        elif "type=\"submit\"" in clean_line:
            note = "6. 주문하기 제출 버튼"
        elif "type=\"reset\"" in clean_line:
            note = "7. 다시 작성 초기화 버튼"
        ws3.append([idx, clean_line, note])

    # -------------------------------------------------------------
    # 스타일 적용 헬퍼
    # -------------------------------------------------------------
    header_fill = PatternFill(start_color="29465B", end_color="29465B", fill_type="solid")
    header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="맑은 고딕", size=10)
    code_font = Font(name="Consolas", size=9)

    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3")
    )
    alt_fill = PatternFill(start_color="F7F9FA", end_color="F7F9FA", fill_type="solid")

    for sheet in [ws1, ws2, ws3]:
        # 헤더 행 서식
        for col_idx in range(1, sheet.max_column + 1):
            cell = sheet.cell(row=2, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        # 데이터 행 서식
        for row_idx in range(3, sheet.max_row + 1):
            sheet.row_dimensions[row_idx].height = 22 if sheet != ws1 else 32
            for col_idx in range(1, sheet.max_column + 1):
                cell = sheet.cell(row=row_idx, column=col_idx)
                cell.font = data_font if sheet != ws3 or col_idx != 2 else code_font
                cell.border = thin_border

                if row_idx % 2 == 0:
                    cell.fill = alt_fill

                # 정렬 규칙
                if col_idx == 1:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif sheet == ws1 and col_idx in [3, 4, 6]:
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                elif sheet == ws2 and col_idx in [1, 2]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        # 열 너비 자동 맞춤
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row == 1:  # 타이틀 제외
                    continue
                val_str = str(cell.value or "")
                for line_str in val_str.split("\n"):
                    # 한글 등 유니코드는 2바이트 계산
                    line_len = sum(2 if ord(c) > 127 else 1 for c in line_str)
                    if line_len > max_len:
                        max_len = line_len
            adjusted_width = min(max(max_len + 4, 12), 65)
            sheet.column_dimensions[col_letter].width = adjusted_width

    # ws3 코드 컬럼 너비 넉넉히
    ws3.column_dimensions["A"].width = 10
    ws3.column_dimensions["B"].width = 100
    ws3.column_dimensions["C"].width = 25

    excel_path = r"c:\Users\dongg\OneDrive\Desktop\AI개발과실전\4주차\카페프로젝트\바이브_카페_주문서_명세서.xlsx"
    wb.save(excel_path)
    print(f"엑셀 파일 저장 완료: {excel_path}")

    # CSV 파일도 추가로 저장 (UTF-8 with BOM)
    csv_path = r"c:\Users\dongg\OneDrive\Desktop\AI개발과실전\4주차\카페프로젝트\바이브_카페_주문서_항목명세.csv"
    with open(csv_path, "w", encoding="utf-8-sig") as f:
        f.write("순번,입력 항목명,HTML 태그/타입,연결 라벨,필수 여부,상세 설명\n")
        f.write('1,이름,<input type="text">,for="customer-name",필수 (required),주문 고객의 성함을 한 줄 텍스트로 입력받는 필드입니다.\n')
        f.write('2,전화번호,<input type="tel">,for="customer-phone",선택,연락처를 입력받는 필드로 모바일에서 숫자 키패드가 나타납니다.\n')
        f.write('3,음료 선택,"<select>, <option>",for="beverage-select",선택,"아메리카노, 카페라떼, 카페모카, 녹차라떼 중 택1 드롭다운 목록입니다."\n')
        f.write('4,사이즈 선택,<input type="radio">,for="size-s/m/l",선택 (택1),"S(작은), M(보통), L(큰) 중 1개를 고르는 단일 선택 라디오 버튼입니다."\n')
        f.write('5,요청사항,<textarea>,for="special-requests",선택,추가 요구사항을 여러 줄로 자유롭게 입력받는 텍스트 박스입니다.\n')
        f.write('6,주문하기,<button type="submit">,-,-,작성된 폼 데이터를 서버로 전송하는 제출 버튼입니다.\n')
        f.write('7,다시 작성,<button type="reset">,-,-,입력된 모든 필드를 초기 상태로 되돌리는 리셋 버튼입니다.\n')
    print(f"CSV 파일 저장 완료: {csv_path}")

if __name__ == "__main__":
    create_excel_specification()

