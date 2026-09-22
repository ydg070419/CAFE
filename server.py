import http.server
import socketserver
import json
import os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PORT = 8000
EXCEL_FILE = "바이브카페_주문내역_누적.xlsx"

def init_excel_file():
    """엑셀 파일이 없으면 헤더와 스타일을 갖춘 새 파일을 생성합니다."""
    if not os.path.exists(EXCEL_FILE):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "카페주문내역"
        ws.views.sheetView[0].showGridLines = True

        # 타이틀
        ws.merge_cells("A1:G1")
        title_cell = ws["A1"]
        title_cell.value = "☕ 바이브 카페 실시간 누적 주문내역"
        title_cell.font = Font(name="맑은 고딕", size=15, bold=True, color="1F497D")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        # 헤더
        headers = ["주문번호", "주문일시", "고객명", "전화번호", "주문음료", "사이즈", "요청사항"]
        ws.append(headers)
        ws.row_dimensions[2].height = 26

        header_fill = PatternFill(start_color="29465B", end_color="29465B", fill_type="solid")
        header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="CCCCCC"),
            right=Side(style="thin", color="CCCCCC"),
            top=Side(style="thin", color="CCCCCC"),
            bottom=Side(style="thin", color="CCCCCC")
        )

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=2, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        # 기본 열 너비
        col_widths = {"A": 10, "B": 22, "C": 14, "D": 18, "E": 16, "F": 12, "G": 40}
        for col_letter, width in col_widths.items():
            ws.column_dimensions[col_letter].width = width

        wb.save(EXCEL_FILE)
        print(f"새 주문내역 엑셀 파일이 생성되었습니다: {EXCEL_FILE}")

def append_order_to_excel(order_data):
    """새 주문 데이터를 엑셀 파일 마지막 행에 추가합니다."""
    init_excel_file()
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb["카페주문내역"]

    next_row = ws.max_row + 1
    order_num = next_row - 2  # 헤더 제외 순번

    row_values = [
        order_num,
        order_data.get("orderDate", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        order_data.get("customerName", ""),
        order_data.get("customerPhone", ""),
        order_data.get("beverage", ""),
        order_data.get("cupSize", ""),
        order_data.get("specialRequests", "")
    ]

    ws.append(row_values)
    ws.row_dimensions[next_row].height = 24

    thin_border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0")
    )
    alt_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    for col_idx in range(1, len(row_values) + 1):
        cell = ws.cell(row=next_row, column=col_idx)
        cell.font = Font(name="맑은 고딕", size=10)
        cell.border = thin_border
        if next_row % 2 == 0:
            cell.fill = alt_fill

        if col_idx in [1, 2, 4, 6]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_idx in [3]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

    wb.save(EXCEL_FILE)
    print(f"[{datetime.now().strftime('%H:%M:%S')}] 새로운 주문이 엑셀에 저장되었습니다! 고객명: {order_data.get('customerName')} / 음료: {order_data.get('beverage')}")

class OrderServerHandler(http.server.SimpleHTTPRequestHandler):
    def end_headers(self):
        # CORS 허용
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        super().end_headers()

    def do_OPTIONS(self):
        self.send_response(200)
        self.end_headers()

    def do_POST(self):
        if self.path == "/api/order":
            content_length = int(self.headers.get("Content-Length", 0))
            post_body = self.rfile.read(content_length)
            try:
                order_data = json.loads(post_body.decode("utf-8"))
                append_order_to_excel(order_data)
                
                self.send_response(200)
                self.send_header("Content-Type", "application/json; charset=utf-8")
                self.end_headers()
                response = {"status": "success", "message": "엑셀 파일에 주문이 성공적으로 저장되었습니다."}
                self.wfile.write(json.dumps(response, ensure_ascii=False).encode("utf-8"))
            except Exception as e:
                self.send_response(500)
                self.send_header("Content-Type", "application/json; charset=utf-8")
                self.end_headers()
                response = {"status": "error", "message": str(e)}
                self.wfile.write(json.dumps(response, ensure_ascii=False).encode("utf-8"))
        else:
            self.send_response(404)
            self.end_headers()

if __name__ == "__main__":
    init_excel_file()
    print("=" * 60)
    print(f"☕ 바이브 카페 주문 서버가 시작되었습니다! (포트: {PORT})")
    print(f"👉 웹 브라우저 접속 주소: http://localhost:{PORT}/index.html")
    print(f"📁 주문 접수 시 [{EXCEL_FILE}]에 자동 누적 저장됩니다.")
    print("=" * 60)
    with socketserver.TCPServer(("", PORT), OrderServerHandler) as httpd:
        try:
            httpd.serve_forever()
        except KeyboardInterrupt:
            print("\n서버가 종료되었습니다.")

